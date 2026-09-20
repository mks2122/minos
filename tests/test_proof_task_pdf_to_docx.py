"""The proof task: convert a PDF to a Word document.

This is the task the whole generality argument was chosen against. It is the
thing a user asks for in one sentence, it has no typed adapter and never will,
and either the file exists at the end or it does not — no judgement call.

Nobody wrote a `doc.convert` operation. Nobody is going to. The model writes a
script, the sandbox runs it, and the result is promoted through the broker as an
ordinary checkpointed, verified, reversible write.

Skipped when the optional sandbox packages are absent (`uv sync --extra
sandbox`), because a missing dependency is a fact about the environment rather
than a failure of the runtime — and `minos doctor` reports it by name.
"""

from __future__ import annotations

import pytest

from minos.audit import AuditLog
from minos.broker import Broker
from minos.checkpoint import FileCheckpointStore
from minos.router import Router
from minos.scopes import ScopeSet
from minos.tiers.l2_code import CodeAdapter
from minos.types import ActionRequest
from minos.undo import find, perform_undo

pytest.importorskip("pypdf", reason="needs: uv sync --extra sandbox")
pytest.importorskip("docx", reason="needs: uv sync --extra sandbox")


# What a model would plausibly write for "convert this PDF to Word".
CONVERT_PDF_TO_DOCX = """
from pathlib import Path
from pypdf import PdfReader
from docx import Document

reader = PdfReader('materials/report.pdf')
document = Document()

for number, page in enumerate(reader.pages, start=1):
    text = (page.extract_text() or '').strip()
    document.add_heading(f'Page {number}', level=2)
    for paragraph in text.split('\\n'):
        if paragraph.strip():
            document.add_paragraph(paragraph.strip())

document.save('out/report.docx')
print(f'converted {len(reader.pages)} page(s)')
"""


def _make_pdf(path):
    """A real PDF, written by the library the sandbox will read it back with.

    Hand-rolling PDF bytes was tried and produced a file pypdf rejects -- which
    is a decent argument for the whole premise of this milestone: use the
    library, do not reimplement the format.
    """
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with open(path, "wb") as handle:
        writer.write(handle)


@pytest.fixture
def rig(tmp_path):
    state = tmp_path / ".minos"
    workspace = tmp_path / "ws"
    workspace.mkdir()
    _make_pdf(workspace / "report.pdf")

    adapter = CodeAdapter(state=state)
    broker = Broker(
        scopes=ScopeSet.parse(
            [
                f"fs.read:{workspace}/**",
                f"fs.write:{workspace}/**",
                f"code.run:{state}/**",
            ]
        ),
        audit=AuditLog(state / "audit.jsonl"),
        store=FileCheckpointStore(state / "checkpoints"),
    )
    return {
        "workspace": workspace,
        "broker": broker,
        "router": Router(adapters=(adapter,)),
    }


def act(rig, operation, **params):
    request = ActionRequest(
        goal_id="pdf2docx", intent=operation, operation=operation, params=params
    )
    routed = rig["router"].route(request)
    return rig["broker"].submit(routed.invocation, routed.execute)


def test_convert_a_pdf_to_a_word_document(rig):
    """The one-sentence request, start to finish."""
    source = rig["workspace"] / "report.pdf"
    destination = rig["workspace"] / "report.docx"

    ran = act(rig, "code.run", code=CONVERT_PDF_TO_DOCX, materials=[str(source)])
    assert ran.status == "ok", ran.error
    assert ran.result.ok, f"{ran.result.stderr}\n{ran.result.detail}"
    assert "converted 1 page" in ran.result.stdout

    promoted = act(rig, "code.materialize", artifact="report.docx", path=str(destination))
    assert promoted.status == "ok", promoted.error

    # It is a real .docx, not a file with the right extension.
    from docx import Document

    assert destination.exists()
    assert Document(str(destination)).paragraphs, "the document has no content"


def test_the_conversion_is_verified_and_undoable(rig):
    """The capability is new; the guarantees are the ones the project always had."""
    source = rig["workspace"] / "report.pdf"
    destination = rig["workspace"] / "report.docx"
    destination.write_text("a file that was already here", encoding="utf-8")

    act(rig, "code.run", code=CONVERT_PDF_TO_DOCX, materials=[str(source)])
    promoted = act(rig, "code.materialize", artifact="report.docx", path=str(destination))

    assert promoted.observed is not None
    assert promoted.observed.verifiable and promoted.observed.matched

    action = find(rig["broker"].audit, rig["broker"].store, "last")
    result, redo = perform_undo(rig["broker"].audit, rig["broker"].store, action)

    assert result.succeeded
    assert destination.read_text(encoding="utf-8") == "a file that was already here"
    assert redo is not None, "the undo is itself undoable"


def test_the_source_pdf_is_never_touched(rig):
    """The script gets a copy. The original is not in play."""
    source = rig["workspace"] / "report.pdf"
    before = source.read_bytes()

    act(rig, "code.run", code=CONVERT_PDF_TO_DOCX, materials=[str(source)])

    assert source.read_bytes() == before


def test_a_missing_package_is_reported_by_name(rig):
    """ "No module named 'x'" is a worse answer than what to install."""
    outcome = act(
        rig,
        "code.run",
        code="import pymupdf",  # in RECOMMENDED, deliberately not installed above
        materials=[],
    )

    assert outcome.status == "ok"
    assert not outcome.result.ok
    assert "pymupdf" in outcome.result.detail
    assert "uv pip install" in outcome.result.detail
